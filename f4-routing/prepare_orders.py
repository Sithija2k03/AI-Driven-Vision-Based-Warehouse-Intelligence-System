"""
F4 Order Preparation (v2) — Fine Enterprises WIS
================================================
Turns the September sales order export into clean pick orders that only
use locations on the finalized warehouse map.

Input : Sep_Sales_Orders_Stock_Report.xlsx (any sheet name)
          sales order lines — one row per order line,
          with Item Code added by XLOOKUP ("Not Found" when it failed)
        item_locations_real.json (from location_mapper.py)

Rules (agreed with Sithija, 2026-10-12)
  * ON-MAP line : item's Stock Report location is on the finalized map.
  * OFF-MAP line: item is in an excluded section, has no location, or
                  cannot be identified.
  * Keep unchanged : no off-map lines.
  * Remove         : every line is off-map, OR off-map lines make up
                     50% or more of the order's lines.
  * Keep reshaped  : otherwise (off-map share below 50%) -> the off-map
                     lines are dropped and the order continues with the
                     remaining on-map lines.
  MAX_OFF_SHARE below sets the 50% threshold.

Other handling
  * Item Code is the join key. For "Not Found" codes the item name is
    matched to the Stock Report (exact, punctuation-insensitive, or unique
    prefix of a truncated name). No fuzzy matching.
  * Dates: Excel turned day<=12 dates into wrong datetimes (1 Sep became
    9 Jan). They are swapped back; text dates are parsed as DD-MM-YYYY.
  * The same item twice in one order is merged (quantities added).
  * Zero-stock items stay routable: they have a slot and were ordered.
  * Customer names (Particulars), salesmen and prices are not carried
    into the outputs.

Outputs: orders_clean.json   orders + on-map lines (input for the RL env)
         orders_clean.csv    one row per kept order line
         orders_removed.csv  one row per removed order, with reason
Run    : python prepare_orders.py
"""

import csv
import datetime as dt
import json
import re
import sys
from collections import Counter, defaultdict

import openpyxl

ORDERS_FILE = 'Sep_Sales_Orders_Stock_Report.xlsx'
ORDERS_SHEET = 'FE_ListofSalesOrderVouchers'
MAX_OFF_SHARE = 0.50     # remove an order when off-map lines are >= this share


def find_sheet(path, preferred, required_headers):
    """Open the sheet named `preferred` (ignoring case/spaces); if it is not
    there, use the first sheet whose header row contains required_headers."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    want = re.sub(r'\s+', '', preferred).lower()
    for ws in wb.worksheets:
        if re.sub(r'\s+', '', ws.title).lower() == want:
            return ws
    for ws in wb.worksheets:
        first = next(ws.iter_rows(max_row=1, values_only=True), ())
        head = {re.sub(r'\s+', ' ', str(h or '')).strip().upper() for h in first}
        if all(h.upper() in head for h in required_headers):
            print(f"  Sheet '{preferred}' not found in {path}; using '{ws.title}' (matching columns)")
            return ws
    raise SystemExit(
        f"\nERROR: {path} has no sheet named '{preferred}' and no sheet with columns "
        f"{required_headers}.\nSheets in this file: {wb.sheetnames}\n"
        f"Check that you copied the right workbook into this folder.")


def norm(s):
    return re.sub(r'\s+', ' ', str(s).replace('"', '').upper()).strip()


def key(s):
    return re.sub(r'[^A-Z0-9]', '', str(s).upper())


def fix_date(v):
    if isinstance(v, dt.datetime):          # Excel read DD-MM as MM-DD
        return dt.date(v.year, v.day, v.month)
    if isinstance(v, str) and v.strip():
        return dt.datetime.strptime(v.strip(), '%d-%m-%Y').date()
    return None


def read_orders(path, sheet):
    ws = find_sheet(path, sheet, ['Vch/Bill No.', 'Item Code', 'Qty.'])
    rows = ws.iter_rows(values_only=True)
    head = [re.sub(r'\s+', ' ', str(h or '')).strip() for h in next(rows)]
    c = {h: i for i, h in enumerate(head)}
    orders = {}
    for r in rows:
        vch = r[c['Vch/Bill No.']]
        if not vch:
            continue
        o = orders.get(vch)
        if o is None:
            o = orders[vch] = {'so_no': vch, 'series': str(vch).split('/')[0], 'date': None,
                               'delivery_type': None, 'lines': []}
        if r[c['Date']] and o['date'] is None:
            o['date'] = fix_date(r[c['Date']])
            o['delivery_type'] = r[c['Type']]
        try:
            qty = float(r[c['Qty.']] or 0)
        except (TypeError, ValueError):
            qty = 0.0
        o['lines'].append({'code_so': str(r[c['Item Code']] or '').strip(),
                           'item_name_so': str(r[c['Item Details']] or '').strip(),
                           'qty': qty, 'unit': r[c['Unit']]})
    return list(orders.values())


def build_name_matcher(items):
    by_norm, by_key = defaultdict(list), defaultdict(list)
    for it in items:
        by_norm[norm(it['item_name'])].append(it)
        by_key[key(it['item_name'])].append(it)
    keys = list(by_key)

    def match(name):
        n, k = norm(name), key(name)
        if len(by_norm.get(n, [])) == 1:
            return by_norm[n][0], 'name exact'
        if len(by_key.get(k, [])) == 1:
            return by_key[k][0], 'name punctuation'
        if len(k) >= 12:
            pre = [x for x in keys if x.startswith(k) and len(by_key[x]) == 1]
            if len(pre) == 1:
                return by_key[pre[0]][0], 'name prefix'
        return None, 'not identified'
    return match


def keep_status(n_on, n_off):
    if n_on and n_off == 0:
        return 'original'
    if not n_on:
        return None
    return 'reshaped' if n_off / (n_on + n_off) < MAX_OFF_SHARE else None


def run(orders_path=ORDERS_FILE, items_path='item_locations_real.json'):
    items = json.load(open(items_path, encoding='utf-8'))['items']
    by_code = {it['sku']: it for it in items}
    match_name = build_name_matcher(items)
    orders = read_orders(orders_path, ORDERS_SHEET)

    how_stats, off_reasons = Counter(), Counter()
    kept, removed = [], []
    for o in orders:
        on, off = [], []
        for ln in o['lines']:
            it = by_code.get(ln['code_so'])
            how = 'item code'
            if it is None:
                it, how = match_name(ln['item_name_so'])
            how_stats[how] += 1
            if it is None:
                off.append({**ln, 'off_reason': 'item not identified'})
            elif not it['routable']:
                off.append({**ln, 'sku': it['sku'], 'off_reason': it['reason']})
            else:
                on.append({'sku': it['sku'], 'item_name': it['item_name'], 'qty': ln['qty'],
                           'unit': ln['unit'], 'location': it['slot_code'],
                           'pick_node': it['pick_node'], 'access_nodes': it['access_nodes'],
                           'floor': it['floor'], 'column': it['column'], 'level': it['level'],
                           'in_stock_end_sep': it['in_stock'], 'matched_by': how})
        for ln in off:
            r = ln['off_reason']
            off_reasons['excluded section' if r.startswith('excluded') else r] += 1

        merged = {}
        for ln in on:
            if ln['sku'] in merged:
                merged[ln['sku']]['qty'] += ln['qty']
            else:
                merged[ln['sku']] = ln
        on = list(merged.values())

        n_on, n_off = len(on), len(off)
        status = keep_status(n_on, n_off)
        base = {'so_no': o['so_no'], 'series': o['series'],
                'date': o['date'].isoformat() if o['date'] else None,
                'delivery_type': o['delivery_type']}
        if status:
            kept.append({**base, 'status': status, 'n_lines': n_on,
                         'dropped_lines': [{'item': l['item_name_so'], 'reason': l['off_reason']} for l in off],
                         'lines': on})
        else:
            why = ('all lines off-map' if n_on == 0 else
                   f'off-map share {n_off / (n_on + n_off):.0%} (>= {MAX_OFF_SHARE:.0%})')
            removed.append({**base, 'on_map_lines': n_on, 'off_map_lines': n_off, 'reason': why})

    kept.sort(key=lambda o: (o['date'] or '', o['so_no']))
    sizes = [o['n_lines'] for o in kept]
    all_sizes = sorted(len(o['lines']) for o in orders)
    floors = Counter('+'.join(sorted({l['floor'] for l in o['lines']})) for o in kept)
    dates = sorted(o['date'] for o in kept if o['date'])
    summary = {
        'source_file': orders_path,
        'period': f'{dates[0]} to {dates[-1]}',
        'orders_in_export': len(orders),
        'orders_by_series': dict(Counter(o['series'] for o in orders)),
        'lines_in_export': sum(all_sizes),
        'lines_per_order_export': {'median': all_sizes[len(all_sizes) // 2],
                                   'mean': round(sum(all_sizes) / len(all_sizes), 1)},
        'item_identification': dict(how_stats),
        'off_map_lines_by_reason': dict(off_reasons),
        'rule': f'remove if all lines off-map or off-map share >= {MAX_OFF_SHARE:.0%}; '
                'otherwise drop off-map lines and keep the order',
        'removed_by_reason': dict(Counter('all lines off-map' if r['on_map_lines'] == 0
                                          else 'off-map share >= 50%' for r in removed)),
        'orders_kept_original': sum(1 for o in kept if o['status'] == 'original'),
        'orders_kept_reshaped': sum(1 for o in kept if o['status'] == 'reshaped'),
        'orders_removed': len(removed),
        'lines_kept': sum(sizes),
        'lines_per_order_kept': {'min': min(sizes), 'median': sorted(sizes)[len(sizes) // 2],
                                 'mean': round(sum(sizes) / len(sizes), 1), 'max': max(sizes)},
        'orders_by_floor': {'ground only': floors.get('G', 0), 'ground + mezzanine': floors.get('G+M', 0),
                            'mezzanine only': floors.get('M', 0)},
        'lines_zero_stock_end_sep': sum(1 for o in kept for l in o['lines'] if not l['in_stock_end_sep']),
        'distinct_skus_ordered': len({l['sku'] for o in kept for l in o['lines']}),
        'orders_per_day': dict(sorted(Counter(o['date'] for o in kept).items())),
    }
    json.dump({'summary': summary, 'orders': kept},
              open('orders_clean.json', 'w', encoding='utf-8'), indent=1, ensure_ascii=False)

    with open('orders_clean.csv', 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['so_no', 'series', 'date', 'delivery_type', 'status', 'sku', 'item_name', 'qty',
                    'unit', 'location', 'pick_node', 'floor', 'column', 'level', 'in_stock_end_sep'])
        for o in kept:
            for l in o['lines']:
                w.writerow([o['so_no'], o['series'], o['date'], o['delivery_type'], o['status'],
                            l['sku'], l['item_name'], l['qty'], l['unit'], l['location'],
                            l['pick_node'], l['floor'], l['column'], l['level'], l['in_stock_end_sep']])
    with open('orders_removed.csv', 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=['so_no', 'series', 'date', 'delivery_type',
                                          'on_map_lines', 'off_map_lines', 'reason'])
        w.writeheader()
        w.writerows(removed)

    s = summary
    print('=' * 64)
    print(f"  Export: {s['orders_in_export']:,} orders, {s['lines_in_export']:,} lines ({s['period']})")
    print(f"  Series: {s['orders_by_series']}")
    print(f"  Item identification: {s['item_identification']}")
    print(f"  Off-map lines: {s['off_map_lines_by_reason']}")
    print(f"  Rule: {s['rule']}")
    print('-' * 64)
    print(f"  Kept unchanged : {s['orders_kept_original']:,}")
    print(f"  Kept reshaped  : {s['orders_kept_reshaped']:,}")
    print(f"  Removed        : {s['orders_removed']:,}  {s['removed_by_reason']}")
    print(f"  Lines kept     : {s['lines_kept']:,}  per order {s['lines_per_order_kept']}")
    print(f"  Export lines/order: {s['lines_per_order_export']}")
    print(f"  Floors: {s['orders_by_floor']}")
    print(f"  Kept lines with zero stock at end of Sep: {s['lines_zero_stock_end_sep']:,}")
    print('=' * 64)
    return summary


if __name__ == '__main__':
    run(*sys.argv[1:])
